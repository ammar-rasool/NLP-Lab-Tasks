# Lab 02 - Tasks 1 to 9
# Simple solutions. Each task and subpart is clearly commented.

import re
import csv
import json
import math
from collections import Counter


# ============================================================
# TASK 1: Multi-Page PDF Extraction with Word Frequency Analysis
# ============================================================
def task1():
    import PyPDF2
    pdf_path = input("\nTask 1 - Enter PDF file path: ")
    try:
        # Task 1(a)
        with open(pdf_path, "rb") as file:
            reader = PyPDF2.PdfReader(file)
            pages = [page.extract_text() or "" for page in reader.pages]
        print("\nTask 1(a)")
        print("Total pages extracted:", len(pages))

        # Task 1(b)
        text = " ".join(pages)
        text = re.sub(r"[\d\W_]+", " ", text.lower())
        tokens = text.split()
        print("\nTask 1(b)")
        print("Total token count:", len(tokens))

        # Task 1(c)
        stopwords = {"the","is","a","an","and","or","of","to","in","on","for","with","as","by","at","from","this","that","it","are","was","were","be","been","but"}
        words = [word for word in tokens if word not in stopwords]
        frequency = Counter(words)
        print("\nTask 1(c)")
        for word, count in frequency.most_common(10):
            print(word, ":", count)

        # Task 1(d)
        keyword = input("\nTask 1(d) - Enter keyword: ").lower()
        found = False
        for page_number, page_text in enumerate(pages, 1):
            for sentence in re.split(r"[.!?]", page_text):
                if keyword in sentence.lower():
                    print("Page", page_number, ":", sentence.strip())
                    found = True
        if not found:
            print("No sentence containing the keyword was found.")
    except Exception as e:
        print("Task 1 error:", e)


# ============================================================
# TASK 2: Structured Page Metadata Extraction and CSV Export
# ============================================================
class PhraseNotFoundError(Exception):
    pass


def task2():
    import PyPDF2
    pdf_path = input("\nTask 2 - Enter PDF file path: ")
    try:
        # Task 2(a)
        with open(pdf_path, "rb") as file:
            reader = PyPDF2.PdfReader(file)
            pages_data = []
            for number, page in enumerate(reader.pages, 1):
                text = page.extract_text() or ""
                pages_data.append({
                    "page_number": number,
                    "raw_text": text,
                    "word_count": len(text.split()),
                    "char_count": len(text.replace(" ", ""))
                })
        print("\nTask 2(a)")
        print(pages_data[:2])

        # Task 2(b)
        non_zero = [p for p in pages_data if p["word_count"] > 0]
        print("\nTask 2(b)")
        if non_zero:
            highest = max(non_zero, key=lambda p: p["word_count"])
            lowest = min(non_zero, key=lambda p: p["word_count"])
            average = sum(p["word_count"] for p in pages_data) / len(pages_data)
            print("Page with highest word count:", highest["page_number"])
            print("Full content:\n", highest["raw_text"])
            print("Page with lowest non-zero word count:", lowest["page_number"])
            print("Average word count:", round(average, 2))
        else:
            print("All pages contain zero words.")

        # Task 2(c)
        phrase = input("\nTask 2(c) - Enter search phrase: ").lower()
        matches = []
        for page in pages_data:
            count = page["raw_text"].lower().count(phrase)
            if count:
                matches.append((page["page_number"], count))
        try:
            if not matches:
                raise PhraseNotFoundError("The phrase was not found on any page.")
            for page_number, count in matches:
                print("Page", page_number, "-", count, "occurrence(s)")
        except PhraseNotFoundError as e:
            print(e)

        # Task 2(d)
        with open("pdf_summary.csv", "w", newline="", encoding="utf-8") as file:
            writer = csv.writer(file)
            writer.writerow(["page_number", "word_count", "char_count"])
            for page in pages_data:
                writer.writerow([page["page_number"], page["word_count"], page["char_count"]])
        print("\nTask 2(d)")
        with open("pdf_summary.csv", "r", encoding="utf-8") as file:
            for row in csv.reader(file):
                print(row)
    except Exception as e:
        print("Task 2 error:", e)


# ============================================================
# TASK 3: PDF Text Preprocessing Pipeline for NLP
# ============================================================
def task3():
    import PyPDF2
    from nltk.corpus import stopwords
    pdf_path = input("\nTask 3 - Enter PDF file path: ")
    try:
        with open(pdf_path, "rb") as file:
            reader = PyPDF2.PdfReader(file)
            text = " ".join((page.extract_text() or "") for page in reader.pages)

        # Task 3(a)
        before = len(text)
        text = re.sub(r"https?://\S+|www\.\S+", " ", text)
        text = re.sub(r"\d+", " ", text)
        text = re.sub(r"[^A-Za-z.\s]", " ", text)
        text = re.sub(r"\s+", " ", text).strip()
        print("\nTask 3(a)")
        print("Character count before cleaning:", before)
        print("Character count after cleaning:", len(text))

        # Task 3(b)
        text = text.lower()
        sentences = re.split(r"[.!?]", text)
        before_empty = len(sentences)
        sentences = [s.strip() for s in sentences]
        sentences = [s for s in sentences if s]
        print("\nTask 3(b)")
        print("Sentence count before removing empty strings:", before_empty)
        print("Sentence count after removing empty strings:", len(sentences))

        # Task 3(c)
        try:
            stop_words = set(stopwords.words("english"))
        except LookupError:
            print("NLTK stopwords data is missing. Run: nltk.download('stopwords')")
            return
        filtered = []
        for sentence in sentences:
            words = [w for w in sentence.split() if w not in stop_words]
            if len(words) >= 5:
                filtered.append(" ".join(words))
        print("\nTask 3(c)")
        print("Sentence count before filter:", len(sentences))
        print("Sentence count after filter:", len(filtered))
        print("Three sample sentences:")
        for sentence in filtered[:3]:
            print(sentence)

        # Task 3(d)
        with open("nlp_ready_output.txt", "w", encoding="utf-8") as file:
            for sentence in filtered:
                file.write(sentence + "\n")
        with open("nlp_ready_output.txt", "r", encoding="utf-8") as file:
            lines = [line.strip() for line in file if line.strip()]
        print("\nTask 3(d)")
        print("Total lines written:", len(lines))
        if lines:
            print("Longest sentence:", max(lines, key=len))
            print("Shortest sentence:", min(lines, key=len))
    except Exception as e:
        print("Task 3 error:", e)


# ============================================================
# TASK 4: DOCX Paragraph Classification and Statistical Analysis
# ============================================================
def task4():
    from docx import Document
    docx_path = input("\nTask 4 - Enter DOCX file path: ")
    try:
        doc = Document(docx_path)

        # Task 4(a)
        paragraphs = [{"text": p.text, "style": p.style.name} for p in doc.paragraphs]
        styles = set(p["style"] for p in paragraphs)
        print("\nTask 4(a)")
        print("Total paragraph count:", len(paragraphs))
        print("Unique style names:", styles)

        # Task 4(b)
        headings, body, other = [], [], []
        for item in paragraphs:
            if item["style"].startswith("Heading"):
                headings.append(item)
            elif item["style"] == "Normal" or "Body" in item["style"]:
                body.append(item)
            else:
                other.append(item)
        print("\nTask 4(b)")
        print("Headings:", len(headings))
        print("Body:", len(body))
        print("Other/Empty:", len(other))

        # Task 4(c)
        qualifying = []
        for item in body:
            words = item["text"].split()
            if len(words) > 5:
                qualifying.append({
                    "text": item["text"],
                    "word_count": len(words),
                    "char_count": len(item["text"].replace(" ", ""))
                })
        print("\nTask 4(c)")
        if qualifying:
            print("Highest word count paragraph:", max(qualifying, key=lambda x: x["word_count"]))
            print("Lowest word count above 5:", min(qualifying, key=lambda x: x["word_count"]))
            print("Mean word count:", round(sum(x["word_count"] for x in qualifying) / len(qualifying), 2))
        else:
            print("No qualifying body paragraphs.")

        # Task 4(d)
        style_dict = {}
        for item in paragraphs:
            style = item["style"]
            if style not in style_dict:
                style_dict[style] = {"count": 0, "texts": []}
            style_dict[style]["count"] += 1
            if len(style_dict[style]["texts"]) < 2:
                style_dict[style]["texts"].append(item["text"])
        style_dict = dict(sorted(style_dict.items(), key=lambda x: x[1]["count"], reverse=True))
        print("\nTask 4(d)")
        print(json.dumps(style_dict, indent=4, ensure_ascii=False))
    except Exception as e:
        print("Task 4 error:", e)


# ============================================================
# TASK 5: DOCX Table Extraction and Pandas Integration
# ============================================================
def task5():
    from docx import Document
    import pandas as pd
    from docx.oxml.ns import qn
    docx_path = input("\nTask 5 - Enter DOCX file path: ")
    try:
        doc = Document(docx_path)

        # Task 5(a)
        tables = []
        for table in doc.tables:
            table_data = [[cell.text for cell in row.cells] for row in table.rows]
            tables.append(table_data)
        print("\nTask 5(a)")
        print("Total tables:", len(tables))
        for i, table in enumerate(tables, 1):
            print("Table", i, "dimensions:", len(table), "x", len(table[0]) if table else 0)
        if not tables:
            print("No table found.")
            return

        # Task 5(b)
        first_table = tables[0]
        df = pd.DataFrame(first_table[1:], columns=first_table[0])
        print("\nTask 5(b)")
        print(df)
        print("Empty/None cells per column:")
        for column in df.columns:
            empty = df[column].isna().sum() + (df[column].astype(str).str.strip() == "").sum()
            print(column, ":", empty)

        # Task 5(c)
        pre_table_paragraphs = []
        for child in doc.element.body.iterchildren():
            if child.tag == qn("w:tbl"):
                break
            if child.tag == qn("w:p"):
                text = "".join(node.text or "" for node in child.iter() if node.tag == qn("w:t"))
                pre_table_paragraphs.append(text)
        print("\nTask 5(c)")
        print("Pre-table paragraph count:", len(pre_table_paragraphs))

        # Task 5(d)
        with open("docx_extracted.txt", "w", encoding="utf-8") as file:
            file.write("--- PRE-TABLE PARAGRAPHS ---\n")
            for paragraph in pre_table_paragraphs:
                file.write(paragraph + "\n")
            file.write("\n--- TABLE DATA ---\n")
            for row in first_table:
                file.write(" | ".join(row) + "\n")
        print("\nTask 5(d)")
        print("Data written to docx_extracted.txt")
    except Exception as e:
        print("Task 5 error:", e)


# ============================================================
# TASK 6: Multi-Document Comparison, Deduplication, Frequency
# ============================================================
def task6():
    from docx import Document
    doc_a_path = input("\nTask 6 - Enter Document A path: ")
    doc_b_path = input("Task 6 - Enter Document B path: ")
    try:
        doc_a = Document(doc_a_path)
        doc_b = Document(doc_b_path)

        # Task 6(a)
        headings_a = [p.text.strip() for p in doc_a.paragraphs if p.style.name.startswith("Heading")]
        headings_b = [p.text.strip() for p in doc_b.paragraphs if p.style.name.startswith("Heading")]
        set_a, set_b = set(headings_a), set(headings_b)
        print("\nTask 6(a)")
        print("Common headings:", set_a & set_b)
        print("Headings only in Document A:", set_a - set_b)
        print("Headings only in Document B:", set_b - set_a)

        # Task 6(b)
        body_a = [p.text for p in doc_a.paragraphs if p.style.name == "Normal"]
        body_b = [p.text for p in doc_b.paragraphs if p.style.name == "Normal"]
        combined = body_a + body_b
        unique_body, seen = [], set()
        for paragraph in combined:
            key = paragraph.strip().lower()
            if key not in seen:
                seen.add(key)
                unique_body.append(paragraph)
        print("\nTask 6(b)")
        print("Count before deduplication:", len(combined))
        print("Count after deduplication:", len(unique_body))

        # Task 6(c)
        stopwords = {"the","is","a","an","and","or","of","to","in","on","for","with","as","by","at","from","this","that","it","are","was","were"}
        all_words = []
        for paragraph in unique_body:
            words = re.sub(r"[^\w\s]", "", paragraph.lower()).split()
            all_words.extend(w for w in words if w not in stopwords)
        frequency = Counter(all_words)
        print("\nTask 6(c)")
        for word, count in frequency.most_common(20):
            print(word, ":", count)

        # Task 6(d)
        lengths = [len(p.split()) for p in unique_body if p.strip()]
        longest = max(unique_body, key=lambda p: len(p.split()), default="")
        print("\nTask 6(d)")
        print("Total headings across both documents:", len(headings_a) + len(headings_b))
        print("Total unique body paragraphs:", len(unique_body))
        print("Unique word types:", len(frequency))
        print("Average paragraph length:", round(sum(lengths) / len(lengths), 2) if lengths else 0)
        print("Longest paragraph:", longest)
    except Exception as e:
        print("Task 6 error:", e)


# ============================================================
# TASK 7: Multi-Endpoint API Data Joining and Analysis
# ============================================================
def task7():
    import requests
    users_url = "https://jsonplaceholder.typicode.com/users"
    posts_url = "https://jsonplaceholder.typicode.com/posts"
    try:
        # Task 7(a)
        users_response = requests.get(users_url)
        posts_response = requests.get(posts_url)
        if users_response.status_code != 200:
            raise ConnectionError("Users status code: " + str(users_response.status_code))
        if posts_response.status_code != 200:
            raise ConnectionError("Posts status code: " + str(posts_response.status_code))
        users = users_response.json()
        posts = posts_response.json()
        print("\nTask 7(a)")
        print("Both requests returned status 200.")

        # Task 7(b)
        joined = {}
        for user in users:
            user_posts = [p for p in posts if p["userId"] == user["id"]]
            joined[user["name"]] = {
                "email": user["email"],
                "post_count": len(user_posts),
                "titles": [p["title"] for p in user_posts]
            }
        print("\nTask 7(b)")
        print(json.dumps(dict(list(joined.items())[:3]), indent=4))

        # Task 7(c)
        top_user = max(users, key=lambda u: len([p for p in posts if p["userId"] == u["id"]]))
        user_posts = [p for p in posts if p["userId"] == top_user["id"]]
        print("\nTask 7(c)")
        print("Full name:", top_user["name"])
        print("Email:", top_user["email"])
        print("Company:", top_user["company"]["name"])
        for i, post in enumerate(user_posts, 1):
            print(i, ".", post["title"])

        # Task 7(d)
        print("\nTask 7(d)")
        for user in users:
            if "synergies" in user["company"]["bs"].lower() or len(user["address"]["city"]) > 8:
                count = len([p for p in posts if p["userId"] == user["id"]])
                print(user["name"], "| City:", user["address"]["city"], "| Post count:", count)
    except Exception as e:
        print("Task 7 error:", e)


# ============================================================
# TASK 8: To-Do Completion Rate Analysis Across Users
# ============================================================
def task8():
    import requests
    todos_url = "https://jsonplaceholder.typicode.com/todos"
    users_url = "https://jsonplaceholder.typicode.com/users"
    try:
        # Task 8(a)
        response = requests.get(todos_url)
        if response.status_code != 200:
            print("Error: Todo API returned status", response.status_code)
            return
        todos = response.json()
        if not isinstance(todos, list):
            print("Error: parsed result is not a list.")
            return
        if not todos:
            print("Error: todo list is empty.")
            return
        print("\nTask 8(a)")
        print("Total todo records:", len(todos))

        # Task 8(b)
        summary = {}
        for todo in todos:
            user_id = todo["userId"]
            if user_id not in summary:
                summary[user_id] = {"total": 0, "completed": 0, "incomplete": 0}
            summary[user_id]["total"] += 1
            if todo["completed"]:
                summary[user_id]["completed"] += 1
            else:
                summary[user_id]["incomplete"] += 1
        print("\nTask 8(b)")
        for user_id in sorted(summary):
            data = summary[user_id]
            check = data["completed"] + data["incomplete"] == data["total"]
            print("User", user_id, ":", data, "| Check:", check)

        # Task 8(c)
        rates = {uid: round(data["completed"] / data["total"] * 100, 1) for uid, data in summary.items()}
        highest = max(rates, key=rates.get)
        lowest = min(rates, key=rates.get)
        print("\nTask 8(c)")
        print("Highest completion rate:", highest, "-", rates[highest], "%")
        print("Lowest completion rate:", lowest, "-", rates[lowest], "%")
        print("Overall average completion rate:", round(sum(rates.values()) / len(rates), 2), "%")

        # Task 8(d)
        users_response = requests.get(users_url)
        if users_response.status_code != 200:
            print("Could not retrieve users.")
            return
        users = users_response.json()
        user_lookup = {u["id"]: u["name"] for u in users}
        sorted_users = sorted(rates.items(), key=lambda x: x[1], reverse=True)
        print("\nTask 8(d)")
        print(f"{'Rank':<6}{'Name':<25}{'Total Tasks':<15}{'Completed':<12}{'Completion Rate %':<20}")
        for rank, (uid, rate) in enumerate(sorted_users, 1):
            data = summary[uid]
            print(f"{rank:<6}{user_lookup[uid]:<25}{data['total']:<15}{data['completed']:<12}{rate:<20}")
    except Exception as e:
        print("Task 8 error:", e)


# ============================================================
# TASK 9: Nested JSON Flattening, DataFrame Operations, Comments
# ============================================================
def task9():
    import requests
    import pandas as pd
    users_url = "https://jsonplaceholder.typicode.com/users"
    posts_url = "https://jsonplaceholder.typicode.com/posts"
    comments_url = "https://jsonplaceholder.typicode.com/comments"
    try:
        # Task 9(a)
        response = requests.get(users_url)
        if response.status_code != 200:
            raise ConnectionError("Users status code: " + str(response.status_code))
        users = response.json()

        def flatten_user(user):
            return {
                "id": user["id"],
                "name": user["name"],
                "username": user["username"],
                "email": user["email"],
                "city": user["address"]["city"],
                "zipcode": user["address"]["zipcode"],
                "lat": user["address"]["geo"]["lat"],
                "lng": user["address"]["geo"]["lng"],
                "company_name": user["company"]["name"],
                "company_catchphrase": user["company"]["catchPhrase"]
            }
        flattened = [flatten_user(user) for user in users]
        print("\nTask 9(a)")
        print(flattened)

        # Task 9(b)
        df = pd.DataFrame(flattened)
        df["lat"] = df["lat"].astype(float)
        df["lng"] = df["lng"].astype(float)
        df["distance_from_origin"] = df.apply(lambda row: math.sqrt(row["lat"] ** 2 + row["lng"] ** 2), axis=1)
        print("\nTask 9(b)")
        print(df.to_string(index=False))
        print("User with greatest distance from origin:")
        print(df.loc[df["distance_from_origin"].idxmax()])

        # Task 9(c)
        northern = df[df["lat"] > 0]["name"].tolist()
        southern = df[df["lat"] < 0]["name"].tolist()
        north_avg = df[df["lat"] > 0]["distance_from_origin"].mean()
        south_avg = df[df["lat"] < 0]["distance_from_origin"].mean()
        print("\nTask 9(c)")
        print("Northern Hemisphere users:", northern)
        print("Southern Hemisphere users:", southern)
        print("Northern average distance:", round(north_avg, 2))
        print("Southern average distance:", round(south_avg, 2))
        if north_avg > south_avg:
            print("Northern Hemisphere has the higher average distance.")
        elif south_avg > north_avg:
            print("Southern Hemisphere has the higher average distance.")
        else:
            print("Both hemispheres have the same average distance.")

        # Task 9(d)
        posts_response = requests.get(posts_url)
        comments_response = requests.get(comments_url)
        if posts_response.status_code != 200:
            raise ConnectionError("Posts status code: " + str(posts_response.status_code))
        if comments_response.status_code != 200:
            raise ConnectionError("Comments status code: " + str(comments_response.status_code))
        posts = posts_response.json()
        comments = comments_response.json()
        comment_counts = Counter(comment["postId"] for comment in comments)
        results = [(post["id"], post["title"][:40], comment_counts[post["id"]]) for post in posts]
        results.sort(key=lambda x: x[2], reverse=True)
        print("\nTask 9(d)")
        print("Top 5 most-commented posts:")
        for item in results[:5]:
            print(item)
        print("Bottom 5 least-commented posts:")
        for item in results[-5:]:
            print(item)
    except Exception as e:
        print("Task 9 error:", e)


# ============================================================
# MAIN PROGRAM
# ============================================================
def main():
    print("LAB 02 - Python Tasks 1 to 9")
    print("Choose a task. File-based tasks will ask for file paths.")
    while True:
        print("\n1. Task 1\n2. Task 2\n3. Task 3\n4. Task 4\n5. Task 5\n6. Task 6\n7. Task 7\n8. Task 8\n9. Task 9\n0. Exit")
        choice = input("Enter choice: ")
        if choice == "1": task1()
        elif choice == "2": task2()
        elif choice == "3": task3()
        elif choice == "4": task4()
        elif choice == "5": task5()
        elif choice == "6": task6()
        elif choice == "7": task7()
        elif choice == "8": task8()
        elif choice == "9": task9()
        elif choice == "0": break
        else: print("Invalid choice.")


if __name__ == "__main__":
    main()
